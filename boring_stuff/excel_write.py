import openpyxl

wb = openpyxl.Workbook()
print(wb.sheetnames)
sheet = wb.active
print(sheet.title)
sheet.title='Spam Bacon Eggs Sheet'
print(wb.sheetnames)

wb.create_sheet()
print(wb.sheetnames)

wb.create_sheet(index=0, title='First sheet')
print(wb.sheetnames)

wb.create_sheet(index=2, title='Middle sheet')
print(wb.sheetnames)

del wb['Middle sheet']
del wb['Sheet']
print(wb.sheetnames)

wb.save('example_copy.xlsx')