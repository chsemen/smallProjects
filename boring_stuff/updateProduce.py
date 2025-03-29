#! python3
# updateProduce.py - Corrects costs in produce sales spreadsheet.

import openpyxl
import pprint

wb =  openpyxl.load_workbook('produceSales.xlsx')
sheet = wb['Sheet']

PRICE_UPDATES = {'Garlic': 3.07,
                 'Celery': 1.19,
                 'Lemon': 1.27}
changes = {}

for rowNum  in range(2, sheet.max_row):
    produceName = sheet.cell(row=rowNum, column=1).value
    if produceName in PRICE_UPDATES:
        sheet.cell(row=rowNum, column=2).value = PRICE_UPDATES[produceName]
        changes.setdefault(produceName, 0)
        changes[produceName] += 1

pprint.pprint(changes)
for key in changes.keys():
    print(f'{key} : {changes[key]}')
wb.save('updatedProduceSales.xlsx')        