import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string
wb = openpyxl.load_workbook('example.xlsx')
print(type(wb))

print(wb.sheetnames)
sheet=wb['Sheet3']
print(sheet)
print(sheet.title)

anotherSheet = wb.active
print(anotherSheet)

sheet=wb['Sheet1']
print(sheet['A1'])
print(sheet['A1'].value)

c=sheet['B1']
print(c.value)
print(f'Cell {c.coordinate} is {c.value}')
print(sheet['C1'].value)

print(sheet.cell(row=1, column=2))
print(sheet.cell(row=1, column=2).value)

for i in range(1,8,2):
    print(i, sheet.cell(row=1,column=2).value)

print(sheet.max_row)    
print(sheet.max_column)    

print(get_column_letter(1))

print(get_column_letter(2))

print(get_column_letter(27))

print(get_column_letter(900))

print(get_column_letter(sheet.max_column))

print(column_index_from_string('A'))
print(column_index_from_string('AA'))

print(tuple(sheet['A1':'C3']))
for rowOfCellObjects in sheet['A1':'C3']:
    for cellObj in rowOfCellObjects:
        print(cellObj.coordinate, cellObj.value)
    print('---END OF ROW---')

print(list(sheet.columns)[1])
for cellObj in list(sheet.columns)[1]:
    print(cellObj.value)

    