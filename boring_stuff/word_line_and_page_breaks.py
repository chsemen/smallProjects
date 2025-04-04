import docx
import docx.enum
import docx.enum.text
import docx.shape
import docx.shared

doc = docx.Document()
doc.add_paragraph('This is on the first page!')
doc.paragraphs[0].runs[0].add_break(docx.enum.text.WD_BREAK.PAGE)
doc.add_paragraph('This is on the second page!')
doc.add_picture('zophie.png', width=docx.shared.Inches(1), height = docx.shared.Cm(4))

doc.save('twoPages.docx')
